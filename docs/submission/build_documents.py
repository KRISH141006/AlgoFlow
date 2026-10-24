from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parent
REPO='https://github.com/KRISH141006/AlgoFlow'
EXTENSION='https://github.com/KRISH141006/dsa_universal_sync'
BLACK=RGBColor(0,0,0)


def make_document(summary=False):
    d=Document();sec=d.sections[0]
    sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.68)
    sec.left_margin=sec.right_margin=Inches(.75)
    for style in d.styles:
        if not hasattr(style,"font"):
            continue
        style.font.name='Times New Roman';style.font.color.rgb=BLACK
        fonts=style._element.get_or_add_rPr().get_or_add_rFonts()
        for attr in list(fonts.attrib):
            del fonts.attrib[attr]
        for attr in ['ascii','hAnsi','eastAsia','cs']:
            fonts.set(qn('w:'+attr),'Times New Roman')
        if hasattr(style,"paragraph_format"):
            style.paragraph_format.space_after=Pt(6)
        if style._element.rPr is not None:
            color=style._element.rPr.find(qn('w:color'))
            if color is not None:
                for attr in ['themeColor','themeTint','themeShade']:
                    color.attrib.pop(qn('w:'+attr),None)
    d.styles['Normal'].font.size=Pt(11 if summary else 12)
    d.styles['Normal'].paragraph_format.line_spacing=1.12
    d.styles['Title'].font.size=Pt(23)
    d.styles['Title'].font.bold=True
    d.styles['Title'].paragraph_format.space_after=Pt(8)
    d.styles['Subtitle'].font.size=Pt(11)
    d.styles['Heading 1'].font.size=Pt(15)
    d.styles['Heading 1'].paragraph_format.space_before=Pt(12)
    d.styles['Heading 1'].paragraph_format.space_after=Pt(8)
    d.styles['Heading 2'].font.size=Pt(11.5)
    d.styles['Heading 2'].paragraph_format.space_before=Pt(9)
    d.styles['Heading 2'].paragraph_format.space_after=Pt(5)
    d.styles['Caption'].font.size=Pt(9)
    d.styles['Caption'].font.italic=True
    d.styles['Caption'].font.bold=False
    for border in list(d.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    d.core_properties.title='AlgoFlow project summary' if summary else 'AlgoFlow project report'
    d.core_properties.author='Krish Sondagar'
    d.core_properties.subject='Personal DSA revision and solution history'
    if not summary:
        p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        r=p.add_run('AlgoFlow project report   |   ');r.font.size=Pt(8)
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
    return d


def p(d,text,style=None):
    return d.add_paragraph(text,style)


def heading(d,text,level=1):
    return d.add_heading(text,level)


def bullets(d,items):
    for text in items:
        para=d.add_paragraph(text,'List Bullet');para.paragraph_format.space_after=Pt(4)


def link(d,label,url):
    para=d.add_paragraph()
    para.add_run(label+' ').bold=True
    rel=para.part.relate_to(url,RT.HYPERLINK,is_external=True)
    hyperlink=OxmlElement('w:hyperlink');hyperlink.set(qn('r:id'),rel)
    run=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'000000');pr.append(color)
    run.append(pr);text=OxmlElement('w:t');text.text=url;run.append(text);hyperlink.append(run);para._p.append(hyperlink)
    return para


def table(d,headers,rows,widths=None,font=10):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=widths or [6.77/len(headers)]*len(headers)
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,h,w in zip(t.rows[0].cells,headers,widths):c.text=h;c.width=Inches(w)
    for row in rows:
        cells=t.add_row().cells
        for c,text,w in zip(cells,row,widths):c.text=str(text);c.width=Inches(w)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for j,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                margin=OxmlElement('w:'+side);margin.set(qn('w:w'),'80');margin.set(qn('w:type'),'dxa');margins.append(margin)
            tcpr.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'FFFFFF');tcpr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(1);para.paragraph_format.line_spacing=1.08
                if j==len(headers)-1 and headers[-1]=='Result':para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size=Pt(font);run.font.color.rgb=BLACK;run.bold=i==0
    p(d,'').paragraph_format.space_after=Pt(1)
    return t


def newpage(d,title):
    d.add_page_break();heading(d,title)


def screenshots(d,title,left,right,caption,explanation):
    newpage(d,title)
    p(d,explanation)
    for idx,name in enumerate([left,right]):
        para=d.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after=Pt(5)
        run=para.add_run();image=run.add_picture(str(ROOT/'screenshots'/name),width=Inches(3.18))
        docpr=image._inline.docPr;docpr.set('descr',caption.split(';')[idx].strip() if ';' in caption else caption)
    p(d,caption,'Caption')


# System design diagram rendered as an original figure rather than copied imagery.
im=Image.new('RGB',(1400,500),'white');draw=ImageDraw.Draw(im)
fontpath='/System/Library/Fonts/Supplemental/Times New Roman.ttf'
font=ImageFont.truetype(fontpath,27);small=ImageFont.truetype(fontpath,20)
boxes=[(30,105,345,230,'Browser interface','React and TypeScript'),(540,105,860,230,'FastAPI application','Routes and validation'),(1050,105,1370,230,'SQLite database','Persistent records')]
for x1,y1,x2,y2,title,sub in boxes:
    draw.rounded_rectangle((x1,y1,x2,y2),radius=0,fill='white',outline='black',width=2)
    draw.text(((x1+x2)/2,y1+32),title,font=font,fill='black',anchor='mt')
    draw.text(((x1+x2)/2,y1+76),sub,font=small,fill='black',anchor='mt')
for x1,x2 in [(345,540),(860,1050)]:
    draw.line((x1,167,x2-12,167),fill='black',width=3)
    draw.polygon([(x2-12,160),(x2,167),(x2-12,174)],fill='black')
draw.text((700,60),'Local application architecture',font=font,fill='black',anchor='mt')
draw.text((700,285),'Universal Sync exports are imported as JSON',font=font,fill='black',anchor='mt')
draw.text((700,338),'Domain services normalize submissions save versions and complete reviews',font=small,fill='black',anchor='mt')
draw.text((700,385),'SQLite write transactions commit related changes together',font=small,fill='black',anchor='mt')
im.save(ROOT/'architecture.png')

summary=make_document(True)
p(summary,'AlgoFlow','Title')
p(summary,'One page project summary','Subtitle')
p(summary,'Submitted by Krish Sondagar   |   15 October 2026')
link(summary,'Application repository',REPO)
link(summary,'Extension repository',EXTENSION)
heading(summary,'Problem statement',2)
p(summary,'Solving a DSA problem once does not ensure that its approach will be remembered. Accepted code, topic notes, and revision dates are often scattered across coding platforms and local files. AlgoFlow brings these records into one local workspace and schedules deliberate revision.')
heading(summary,'Objectives',2)
p(summary,'Organize accepted LeetCode and GeeksforGeeks solutions; preserve exact code and problem descriptions; schedule problem and theory reviews; record improvements as solution versions; and support safe backup and restoration.')
heading(summary,'Key features',2)
p(summary,'AlgoFlow uses my custom DSA Universal Sync extension to collect code and problem details from LeetCode and GeeksforGeeks. The extension can fetch solved problems together as a batch and automatically save code when a new submission is accepted. Its JSON exports can be imported into AlgoFlow. Other features include a searchable library, duplicate protection, saved solution versions, scheduled reviews, progress tracking, and backup and restore.')
heading(summary,'Technologies and tools',2)
p(summary,'Python, FastAPI, SQLite, React, TypeScript, Vite, JavaScript, browser extension APIs, Pydantic, Bleach, pytest, and GitHub.')
heading(summary,'Implementation',2)
p(summary,'The extension collects solutions and exports them as JSON. The React interface sends imported data and form entries to the FastAPI backend. The backend validates the data and stores problems, topics, solution versions, and review history in SQLite. Reviews follow intervals of 1, 7, 30, and 180 days, and saving a new version keeps earlier code available.')
heading(summary,'Achieved outcome',2)
p(summary,'AlgoFlow imports JSON files generated by the custom extension and organizes the problem details and submitted code in a searchable library. The dashboard shows due reviews and a summary of progress. The application schedules newly imported problems for revision and updates the next review date when a review is completed. Users can revisit saved solutions, record improved versions, and back up their data. This provides a clear workflow for collecting solutions and revising them regularly.')
summary.save(ROOT/'AlgoFlow_One_Page_Summary.docx')

report=make_document()
p(report,'AlgoFlow Project Report','Title')
p(report,'Personal DSA revision and solution history','Subtitle')
p(report,'Submitted by Krish Sondagar   |   15 October 2026')
link(report,'Source repository',REPO)
heading(report,'1 Project overview')
p(report,'AlgoFlow is a local web application that stores accepted DSA solutions, schedules revision, and records how a solution changes over time. This report describes the current application, its system design, database schema, implementation, screenshots, and test results.')
heading(report,'Problem statement',2)
p(report,'Coding platforms record submissions, but the student still needs a practical way to revisit an approach, locate an earlier solution, and distinguish a saved attempt from an accepted submission. Local code files and topic notes do not provide a consistent revision schedule. Repeated imports can also create duplicate records if problem and submission identity are not handled explicitly.')
heading(report,'Objectives',2)
bullets(report,['Consolidate exported LeetCode and GeeksforGeeks solutions in a searchable library.','Preserve exact source code, problem descriptions, topic tags, and solution assessments.','Schedule explicit problem and theory reviews without altering existing revision history.','Provide persistent storage, duplicate protection, and complete backup and additive restore.'])
heading(report,'Key features',2)
p(report,'My custom DSA Universal Sync extension collects solved problems and code from LeetCode and GeeksforGeeks. It supports fetching solved problems together and saving newly accepted submissions automatically as JSON files. AlgoFlow imports these files into a searchable collection. The application provides four views: Today, Library, Problem Detail, and Progress. Users can import JSON, add problems and topics, view stored solutions, save a new version, complete a due review, and export or restore a backup. Acceptance status and review status are separate: an accepted problem can still be pending revision.')
heading(report,'Implementation scope',2)
p(report,'I chose this project to organize my DSA practice and make revision easier. The application uses FastAPI, React, and a local SQLite database. I wrote the Universal Sync extension to collect my solutions, and the web application was developed with AI assistance. The current workflow uses JSON files to transfer data from the extension to AlgoFlow. It is designed for one local user.')

newpage(report,'2 System design')
p(report,'The system has two main parts: the browser extension that collects solutions and the web application that organizes them. The frontend and backend are served from one local address. HTTP requests enter FastAPI routes, pass through request validation and service functions, and read or update the selected SQLite database. The original tracker is read only during preparation.')
para=report.add_paragraph();para.add_run().add_picture(str(ROOT/'architecture.png'),width=Inches(6.6))
p(report,'Figure 1 Local application architecture','Caption')
heading(report,'Component responsibilities',2)
table(report,['Component','Responsibility'],[
('Universal Sync extension','Collects solved problems in batches and captures newly accepted code as JSON.'),
('React interface','Displays queues, filters, descriptions, code snapshots, forms, and progress.'),
('FastAPI routes','Expose /api/v1 endpoints and translate HTTP requests into application operations.'),
('Pydantic schemas','Validate required values, lengths, difficulty, problem URLs, and expected review stage.'),
('Domain services','Normalize platform exports, deduplicate submissions, create snapshots, and apply revision rules.'),
('SQLite access','Use parameterized queries, foreign keys, repeatable migrations, and write transactions.'),
('Backup service','Export all five tables and restore with validation and parent ID remapping.')],[1.65,5.12])
heading(report,'Data flow',2)
p(report,'The extension exports problem details and code to JSON. The user imports the file through AlgoFlow, then each submission is normalized into a common internal format. Valid records are stored together in a transaction. The backend returns counts of created, updated, skipped, and invalid records. The interface refreshes its library and statistics from the database after the operation.')
heading(report,'Data preservation',2)
p(report,'The original tracker is backed up before its records are copied into the application. Migration preserves existing problem labels, review stages, dates, and topic history. Backup export and restoration provide a way to recover records without losing saved solution versions.')

newpage(report,'3 Database schema for problems and topics')
p(report,'SQLite stores five related tables. Integer primary keys identify records inside AlgoFlow; platform IDs and URL slugs identify external problems. JSON tags are stored as TEXT and decoded for API responses.')
heading(report,'Problems table',2)
table(report,['Fields','SQLite type','Purpose'],[
('id','INTEGER primary key','Local problem identifier.'),('name difficulty link','TEXT','Display label, Easy/Medium/Hard classification, and reference URL.'),('date_added','TEXT','Original tracker addition date or administrative creation date.'),('stage status next_review_date','INTEGER and TEXT','Revision stage 0 to 4, pending/mastered state, and due date.'),('platform slug platform_id','TEXT','Platform identity and external metadata.'),('topic_tags company_tags','TEXT containing JSON','Topic and company labels.'),('problem_html','TEXT','Sanitized question description for local reading.'),('imported_at solved_at','TEXT nullable','Import timestamp; solve timestamp is left unavailable when absent.'),('historical_import','INTEGER','Distinguishes historical collection imports from manual additions for activity counting.'),('solve_state','TEXT','Accepted or unverified solution state.')],[2.25,1.6,2.92])
heading(report,'Topics table',2)
table(report,['Fields','SQLite type','Purpose'],[
('id','INTEGER primary key','Local theory topic identifier.'),('name category notes link','TEXT','Concept name, grouping, personal explanation, and reference.'),('date_added','TEXT','Date the topic was added.'),('stage status next_review_date','INTEGER and TEXT','Theory revision progression and due date.')],[2.25,1.6,2.92])
heading(report,'Identity and migration rules',2)
p(report,'Problems are matched by platform and canonical URL slug. Existing labels, IDs, stages, and dates are retained when an import enriches a tracker record. Schema upgrades add columns and tables rather than deleting the original records. The problems identity index supports lookup; ingestion serializes writes to prevent competing duplicate inserts.')

newpage(report,'4 Database schema for versions and reviews')
heading(report,'Problem versions table',2)
table(report,['Fields','SQLite type','Purpose'],[
('id problem_id','INTEGER','Snapshot identifier and foreign key to problems.'),('version_number','INTEGER','Sequential number within one problem.'),('commit_message code language','TEXT','Explanation, exact submitted code, and programming language.'),('what_is_good what_is_bad','TEXT','Self-assessment of the approach and improvements.'),('time_complexity space_complexity','TEXT','Optional time and space complexity notes.'),('created_at accepted source','TEXT and INTEGER','Snapshot creation time, acceptance flag, and import/manual origin.'),('submission_id','TEXT nullable','External submission identity when available.'),('fingerprint delivery_key','TEXT','Code digest and repeat-delivery identity.')],[2.35,1.35,3.07])
p(report,'The database enforces unique (problem_id, version_number) and (problem_id, delivery_key) pairs. A foreign key associates every snapshot with a problem. There is no endpoint that edits a saved code snapshot.')
heading(report,'Review history tables',2)
table(report,['Field','Type','Meaning'],[
('id','INTEGER primary key','Review event identifier.'),('problem_id or topic_id','INTEGER foreign key','Links to problems or topics respectively.'),('review_date','TEXT','Completion date in Asia/Kolkata.'),('stage_completed','INTEGER','Stage reached after completing the review.'),('was_on_time','INTEGER','1 when completed on its scheduled day, otherwise 0.')],[2.25,1.6,2.92])
p(report,'review_history belongs to problems; topic_review_history belongs to topics. One problem can have many versions and reviews. One theory topic can have many reviews. Versions and reviews are independent records, so saving code does not silently complete a review.')
heading(report,'Consistency guarantees',2)
p(report,'Every connection enables foreign keys. Related writes run in a transaction using BEGIN IMMEDIATE. Unexpected database errors roll back the operation. Backup restoration validates parent references and remaps identifiers when the receiving database already contains different records.')

newpage(report,'5 Implementation details')
heading(report,'Submission import',2)
p(report,'The normalizer accepts a single object or an array. LeetCode exports use userCode and problem_html; older GFG exports can use submission.code and problem_question. Names, platform identity, slugs, difficulty, code, language, and tag lists are validated. Ambiguous platform information is reported rather than guessed.')
p(report,'A supplied submission ID is used for duplicate detection. Without one, the backend hashes the language and exact code within the problem. Distinct submission IDs preserve separate submissions. Deliberate manual commits may keep identical code with a different assessment.')
heading(report,'Review completion',2)
table(report,['State after action','Next review','Task'],[
('Added at stage 0','Tomorrow','First recall'),('Stage 1 completed','7 days later','Second recall'),('Stage 2 completed','30 days later','Re-solve'),('Stage 3 completed','180 days later','Long-term re-solve'),('Stage 4 completed','None','Mastered')],[2.1,1.8,2.87])
p(report,'A completion request includes expected_stage. The service checks that the item exists, is due, is not mastered, and still matches that stage. It logs the review and advances exactly one stage in the same transaction. A repeated or stale request receives a conflict response.')
heading(report,'Forms and validation',2)
p(report,'React forms support problem entry, theory notes, solution versions, and file uploads. Required fields and URL controls provide browser feedback; Pydantic performs backend validation. Missing or invalid values produce a validation response. Saving a manual version does not establish platform acceptance.')
heading(report,'Backend interfaces',2)
p(report,'The /api/v1 endpoints cover problems, due queues, topics, version lists and commits, explicit review completion, JSON ingestion, statistics, heatmaps, and backup export/restore. FastAPI provides interactive documentation at /docs. The interface and API use the same local origin.')
heading(report,'Security and user management',2)
p(report,'Descriptions are sanitized with Bleach, problem links accept HTTP/HTTPS, and SQL values are parameterized. Uploads are limited to 25 MB and batches to 10,000 records. The server binds to loopback. This release has no authentication or multi-user isolation; it is intended for one local user.')

newpage(report,'6 Testing results and achieved outcome')
p(report,'The automated suite was rerun on 15 October 2026. All 13 tests passed in approximately 3 seconds. Tests use temporary databases. The real export and read-only tracker are used for integration checks when those files are available.')
table(report,['Test scenario','Verified behavior','Result'],[
('Real batch and repeat import','274 exact code snapshots; repeat import skips all records.','Pass'),('Original tracker migration','Original IDs, labels, dates, stages, and history retained.','Pass'),('GFG export variants','Nested and universal formats; distinct submission IDs retained.','Pass'),('Invalid records and HTML','Invalid JSON reported; executable HTML removed.','Pass'),('Review cycle and replay','7/30/180-day intervals; stale completion rejected.','Pass'),('Concurrent review requests','Exactly one request advances the stage.','Pass'),('Version persistence','Manual code survives restart without a review change.','Pass'),('Backup round trip','Versions and reviews restored with remapped parent IDs.','Pass'),('Invalid backup','Orphan references rejected without partial changes.','Pass'),('Import write failure','Simulated failure rolls back the entire valid import.','Pass'),('Manual entries and topics','Required values validated; future reviews rejected.','Pass'),('Identity and URL checks','Platform mismatch and unsafe manual link rejected.','Pass'),('Existing activity enrichment','Import enrichment preserves manual addition activity.','Pass')],[2.25,3.92,.6],font=9.5)
heading(report,'Browser verification',2)
p(report,'Browser checks confirmed search, imported descriptions and code, manual version saving, explicit due-review completion, repeat upload of the 274-record batch, a real GFG import, backup download and additive restore. Restarting the server preserved the saved version and completed review. The frontend TypeScript build passed.')
heading(report,'Achieved outcome',2)
p(report,'AlgoFlow imports JSON files generated by the custom extension and organizes the problem details and submitted code in a searchable library. The dashboard shows due reviews and a summary of progress. The application schedules newly imported problems for revision and updates the next review date when a review is completed. Users can revisit saved solutions, record improved versions, and back up their data. This provides a clear workflow for collecting solutions and revising them regularly.')

screenshots(report,'7 Dashboard and searchable library','01-dashboard.jpg','02-library.jpg','Figure 2 Today dashboard with one due item and 274 accepted solutions; Figure 3 Library filtered to Valid Parentheses with difficulty platform and topic controls.','The dashboard separates due reviews, accepted solutions, completed reviews, and streaks. The library distinguishes accepted and unverified records and provides filters for locating a stored solution. The screenshots show the working application in a browser.')
screenshots(report,'8 Problem description and saved code','03-description.jpg','04-solution.jpg','Figure 4 Sanitized Valid Parentheses description and examples; Figure 5 Exact imported C++ code and its initial solution version.','Problem Detail provides an offline question reader and version selector. The accepted snapshot retains exact code from the export. Import dates are shown separately from solve dates, which remain unavailable when the export contains no solve timestamp.')
screenshots(report,'9 Solution form and progress view','05-version-form.jpg','06-progress.jpg','Figure 6 Form for saving a manual version with message and assessment; Figure 7 Collection statistics and revision activity heatmap.','The version form accepts code, language, a commit message, self-assessment, and complexity notes. The progress view summarizes the collection and actual recorded activity. Historical batch imports do not create synthetic solve events or review activity.')

newpage(report,'10 Conclusion and future scope')
heading(report,'Conclusion',2)
p(report,'AlgoFlow addresses a problem I face during DSA practice: solved questions and saved code are easy to collect but difficult to revise regularly. The application brings problem descriptions, source code, personal notes, and review dates into one place. Earlier solutions remain available when a new version is saved, which makes it easier to compare my approach over time.')
p(report,'The custom Universal Sync extension supports the collection process by fetching solved problems in batches and capturing code after accepted submissions on LeetCode and GeeksforGeeks. These records are exported as JSON and imported into AlgoFlow. Together, the extension and application reduce the manual work needed to maintain a personal solution collection.')
p(report,'The application was checked using the 274-problem export and the existing tracker records. Testing confirmed duplicate protection, review progression, persistent solution versions, and backup restoration. The project meets its main objective of providing a practical workspace for organizing solutions and planning revision.')
heading(report,'Current limitations',2)
p(report,'Data transfer from the extension to AlgoFlow currently requires importing a JSON file. The application does not run or judge saved code, and a manually saved version is not marked as accepted by a coding platform. It stores data locally and does not yet include user accounts or synchronization between devices. Historical records without solve timestamps are kept without assigning an invented solve date.')
heading(report,'Future scope',2)
bullets(report,['Connect the extension directly to AlgoFlow, with a retry queue for submissions that cannot be delivered immediately.','Add a comparison view to highlight changes between saved solution versions.','Support selected local code files so that practice drafts can be stored alongside accepted solutions.','Add accounts and cloud synchronization if the application is extended for use across devices.'])
heading(report,'Project repositories',2)
link(report,'AlgoFlow application',REPO)
link(report,'DSA Universal Sync extension',EXTENSION)
p(report,'The application repository contains the backend, frontend, database migration logic, and automated tests. The extension repository contains the browser scripts used for bulk collection and accepted-submission capture.')
report.save(ROOT/'AlgoFlow_Project_Report.docx')
print('Created summary and report')
